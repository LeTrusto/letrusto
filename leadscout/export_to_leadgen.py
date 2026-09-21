"""Bridge: convert a leadscout leads-*.xlsx into a CSV leadgen's Import page can read.

Only rows with a public business email are exported, since leadgen's outreach
queue requires a valid recipient. The 'Observation' column carries over the
factual social-proof note detected during crawling (never fabricated) so
leadgen's qualification and personalization checks are satisfied honestly.
Rows without a detected signal are still exported but will need a manual
factual observation added on the lead's page in leadgen before they can be
queued for sending -- this is leadgen's built-in anti-spam safeguard and is
intentionally not bypassed here.

Usage:
    python export_to_leadgen.py leads-beauty.xlsx --niche "beauty" --country "" --output leadgen_import.csv
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


def convert(input_path: Path, output_path: Path, niche: str, country: str) -> tuple[int, int]:
    workbook = load_workbook(input_path)
    sheet = workbook.active
    header = [str(cell.value) for cell in sheet[1]]
    index = {name: position for position, name in enumerate(header)}

    exported = 0
    with_observation = 0
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Company", "Website", "Email", "Country", "Niche", "LinkedIn URL", "Observation"])
        for row in sheet.iter_rows(min_row=2, values_only=True):
            email = str(row[index["business_email"]] or "").strip()
            if not email:
                continue
            observation = str(row[index["social_proof_note"]] or "").strip() if "social_proof_note" in index else ""
            writer.writerow([
                row[index["company_name"]] or "",
                row[index["website"]] or "",
                email,
                country,
                niche,
                row[index["linkedin_url"]] or "" if "linkedin_url" in index else "",
                observation,
            ])
            exported += 1
            if observation:
                with_observation += 1
    return exported, with_observation


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert leadscout xlsx output into a leadgen-importable CSV.")
    parser.add_argument("input", help="Path to a leadscout leads-*.xlsx file")
    parser.add_argument("--output", default="leadgen_import.csv", help="CSV file to write")
    parser.add_argument("--niche", default="", help="Niche label to assign, e.g. 'beauty'")
    parser.add_argument("--country", default="", help="Country label to assign")
    args = parser.parse_args()

    exported, with_observation = convert(Path(args.input), Path(args.output), args.niche, args.country)
    print(f"Wrote {exported} row(s) with a public email to {args.output}.")
    print(f"{with_observation} row(s) already have a factual observation detected by the crawl.")
    print(f"{exported - with_observation} row(s) will need a manual factual observation added in leadgen before they can be queued for sending.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
